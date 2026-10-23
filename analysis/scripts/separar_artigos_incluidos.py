#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Exporta apenas os artigos marcados como 'include' na aba FILA_236.

Uso (na raiz do repositório):
  python -m pip install openpyxl
  python analysis/scripts/separar_artigos_incluidos.py \
    "screening/day10/fila_texto_completo_dia10_236_COM_EXTRACAO.xlsm" \
    -o "extraction/Artigos_Incluidos_RSL_RAG_124.xlsx"

A origem (.xlsm) é aberta apenas para leitura e não é alterada.
A saída é um novo arquivo .xlsx, sem macros, com as 29 colunas A:AC,
LINHA ORIGINAL e ALERTA DE CONSISTÊNCIA.
"""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
import re
import sys
import unicodedata

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit(
        "ERRO: falta a biblioteca 'openpyxl'.\n"
        "Instale no mesmo Python com: python -m pip install openpyxl\n"
        "Depois execute novamente o script."
    )

ABA_ORIGEM = "FILA_236"
LINHA_CABECALHO = 3
PRIMEIRA_LINHA_DADOS = 4
NUM_COLUNAS_ORIGINAIS = 29  # A:AC


def normalizar(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "").strip().casefold())
    return "".join(letra for letra in texto if not unicodedata.combining(letra))


def avaliar_alerta(cabecalhos: list[object], registro: list[object]) -> str:
    campos = {
        normalizar(cab): str(registro[i] if registro[i] is not None else "").strip()
        for i, cab in enumerate(cabecalhos)
    }
    alertas: list[str] = []

    for cab in (
        "RAG OPERACIONAL (CI1)",
        "AMBIENTAL SUBSTANTIVO (CI2)",
        "ESTUDO PRIMÁRIO (CI3)",
        "AVALIAÇÃO EMPÍRICA (CI4)",
    ):
        resposta = normalizar(campos.get(normalizar(cab), ""))
        if resposta in {"nao", "no", "false"}:
            alertas.append(f"{cab}: NÃO")
        elif resposta in {"", "nr", "indefinido", "pendente"}:
            alertas.append(f"{cab}: verificar")

    codigo = campos.get(normalizar("CE PRINCIPAL"), "")
    if re.match(r"^CE[1-7]\b", codigo, flags=re.I):
        alertas.append(f"CE principal preenchido: {codigo}")

    return "; ".join(alertas) if alertas else "OK"


def filtrar(origem: Path, destino: Path) -> None:
    if not origem.is_file():
        raise FileNotFoundError(f"Planilha de origem não encontrada: {origem}")
    if origem.resolve() == destino.resolve():
        raise ValueError("O destino não pode ser o mesmo arquivo da origem.")
    if destino.suffix.lower() != ".xlsx":
        raise ValueError("A saída deve ser um arquivo .xlsx")

    entrada = load_workbook(origem, read_only=True, data_only=True)
    try:
        if ABA_ORIGEM not in entrada.sheetnames:
            raise ValueError(
                f"Aba {ABA_ORIGEM!r} não encontrada. Abas: {entrada.sheetnames}"
            )
        aba = entrada[ABA_ORIGEM]
        cabecalhos = [aba.cell(LINHA_CABECALHO, col).value
                      for col in range(1, NUM_COLUNAS_ORIGINAIS + 1)]
        indices = {normalizar(cab): i for i, cab in enumerate(cabecalhos)}
        chave_decisao = normalizar("DECISÃO TEXTO COMPLETO")
        if chave_decisao not in indices:
            raise ValueError(
                "Coluna 'DECISÃO TEXTO COMPLETO' não encontrada na linha 3. "
                "Confira a planilha e a aba FILA_236."
            )
        indice_decisao = indices[chave_decisao]
        indice_id = indices.get(normalizar("ID"), 1)

        registros: list[tuple[list[object], int, str]] = []
        contagens: Counter[str] = Counter()
        ids: set[str] = set()

        for linha_excel, celulas in enumerate(
            aba.iter_rows(min_row=PRIMEIRA_LINHA_DADOS,
                          max_col=NUM_COLUNAS_ORIGINAIS, values_only=True),
            start=PRIMEIRA_LINHA_DADOS,
        ):
            registro = list(celulas)
            if not any(v is not None and str(v).strip() for v in registro):
                continue
            decisao = normalizar(registro[indice_decisao])
            contagens[decisao or "(vazio)"] += 1
            if decisao != "include":
                continue

            codigo = str(registro[indice_id] or "").strip()
            if not codigo:
                raise ValueError(f"Estudo incluído sem ID na linha {linha_excel}")
            if codigo in ids:
                raise ValueError(f"ID duplicado entre os incluídos: {codigo}")
            ids.add(codigo)
            alerta = avaliar_alerta(cabecalhos, registro)
            registros.append((registro, linha_excel, alerta))
    finally:
        entrada.close()

    if not registros:
        raise ValueError("Não há estudos marcados como 'include' nesta planilha.")

    saida = Workbook()
    plan = saida.active
    plan.title = "INCLUIDOS"
    plan.append(cabecalhos + ["LINHA ORIGINAL (FILA_236)", "ALERTA DE CONSISTÊNCIA"])
    alertas: list[str] = []

    for registro, linha_excel, alerta in registros:
        plan.append(registro + [linha_excel, alerta])
        if alerta != "OK":
            codigo = str(registro[indice_id])
            alertas.append(f"{codigo} (linha {linha_excel}): {alerta}")
            plan.cell(plan.max_row, 31).fill = PatternFill(
                "solid", fgColor="FDE8D3"
            )
            plan.cell(plan.max_row, 31).font = Font(color="9A3412", bold=True)

    cab_fill = PatternFill("solid", fgColor="163A4A")
    for celula in plan[1]:
        celula.fill = cab_fill
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    plan.row_dimensions[1].height = 38
    plan.freeze_panes = "A2"
    plan.auto_filter.ref = f"A1:AE{plan.max_row}"

    larguras = {
        "A": 10, "B": 13, "C": 15, "D": 55, "E": 65,
        "F": 22, "G": 27, "H": 50, "I": 19, "J": 22,
        "K": 24, "L": 22, "M": 23, "N": 26, "O": 17,
        "P": 60, "Q": 60, "R": 15, "S": 17, "T": 22,
        "U": 65, "V": 60, "W": 60, "X": 60, "Y": 60,
        "Z": 60, "AA": 60, "AB": 60, "AC": 55,
        "AD": 25, "AE": 55,
    }
    for coluna, largura in larguras.items():
        plan.column_dimensions[coluna].width = largura

    destino.parent.mkdir(parents=True, exist_ok=True)
    saida.save(destino)

    print(f"\nOrigem: {origem}")
    print(f"Registros encontrados: {sum(contagens.values())} — {dict(contagens)}")
    print(f"Incluídos exportados: {len(registros)}")
    print(f"Destino: {destino.resolve()}")
    if alertas:
        print("\nATENÇÃO: inconsistências a revisar (nenhuma decisão foi alterada):")
        for alerta in alertas:
            print(f"  - {alerta}")
    print("\nA planilha de origem não foi modificada.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Exporta somente os estudos 'include' de FILA_236.")
    parser.add_argument("origem", type=Path, help="Arquivo Excel de origem (.xlsm/.xlsx)")
    parser.add_argument("-o", "--saida", type=Path,
                        help="Arquivo .xlsx de saída (opcional)")
    args = parser.parse_args()
    destino = args.saida or args.origem.with_name("artigos_incluidos_filtrados.xlsx")
    try:
        filtrar(args.origem, destino)
    except (OSError, ValueError, PermissionError) as erro:
        sys.exit(f"ERRO: {erro}")


if __name__ == "__main__":
    main()
